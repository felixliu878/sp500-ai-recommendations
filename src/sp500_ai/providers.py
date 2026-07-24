from __future__ import annotations

import io
import hashlib
import json
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime, time as dt_time, timedelta, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


USER_AGENT = "sp500-ai-academic-research/0.1 contact=research-project"
YAHOO_RESEARCH_SCHEMA_VERSION = "yahoo-research-v2"

YAHOO_TO_GICS_SECTOR = {
    "basic materials": "Materials",
    "communication services": "Communication Services",
    "consumer cyclical": "Consumer Discretionary",
    "consumer defensive": "Consumer Staples",
    "energy": "Energy",
    "financial services": "Financials",
    "healthcare": "Health Care",
    "industrials": "Industrials",
    "real estate": "Real Estate",
    "technology": "Information Technology",
    "utilities": "Utilities",
}


def normalize_yahoo_sector(value: Any) -> str | None:
    normalized = " ".join(str(value or "").strip().lower().split())
    return YAHOO_TO_GICS_SECTOR.get(normalized)


def _redact_url_credentials(url: str) -> str:
    return re.sub(
        r"(?i)([?&](?:api[_-]?key|apikey|access[_-]?token|token|key)=)[^&\s]+",
        r"\1[REDACTED]",
        url,
    )


def _request_bytes(url: str, retries: int = 3, timeout: int = 30) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read()
        except (urllib.error.URLError, TimeoutError) as exc:
            last_error = exc
            if attempt + 1 < retries:
                time.sleep(2**attempt)
    raise RuntimeError(
        f"Failed to retrieve {_redact_url_credentials(url)}: {last_error}"
    )


def _request_json(url: str, retries: int = 3) -> Any:
    return json.loads(_request_bytes(url, retries=retries).decode("utf-8"))


def ticker_for_yahoo(symbol: str) -> str:
    return symbol.strip().upper().replace(".", "-")


@dataclass
class PriceResult:
    frame: pd.DataFrame
    metadata: dict[str, Any]
    raw: dict[str, Any]
    provider: str


@dataclass
class ResearchResult:
    fundamentals: dict[str, Any]
    analyst_consensus: dict[str, Any]
    metadata: dict[str, Any]
    raw: dict[str, Any]
    provider: str


class StateStreetUniverseProvider:
    """Download a current, dated SPY holdings snapshot from State Street."""

    def __init__(self, url: str):
        self.url = url

    def fetch(self) -> tuple[pd.DataFrame, dict[str, Any], bytes]:
        payload = _request_bytes(self.url)
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        as_of = str(rows[2][1]) if len(rows) > 2 else ""
        header_index = next(
            i for i, row in enumerate(rows) if row and str(row[0]).strip() == "Name"
        )
        headers = [str(value).strip() if value is not None else "" for value in rows[header_index]]
        records: list[dict[str, Any]] = []
        for row in rows[header_index + 1 :]:
            name = row[0] if row else None
            ticker = row[1] if len(row) > 1 else None
            if not name or not ticker:
                continue
            if str(ticker).upper() in {"USD", "CASH_USD"}:
                continue
            if str(name).upper().startswith("CONTRA "):
                # Temporary corporate-action receivables can appear in fund holdings
                # but are not S&P 500 constituent securities.
                continue
            record = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
            records.append(
                {
                    "symbol": str(record.get("Ticker", "")).strip(),
                    "company": str(record.get("Name", "")).strip(),
                    "identifier": str(record.get("Identifier", "")).strip(),
                    "sedol": str(record.get("SEDOL", "")).strip(),
                    "weight": pd.to_numeric(record.get("Weight"), errors="coerce"),
                    "sector": "" if record.get("Sector") in (None, "-") else str(record.get("Sector")),
                    "shares_held": pd.to_numeric(record.get("Shares Held"), errors="coerce"),
                    "currency": str(record.get("Local Currency", "")).strip(),
                }
            )
        frame = pd.DataFrame.from_records(records)
        metadata = {
            "provider": "State Street SPY daily holdings",
            "source_url": self.url,
            "as_of_label": as_of,
            "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
            "security_count": int(len(frame)),
        }
        return frame, metadata, payload


class YahooChartProvider:
    """Unofficial no-key Yahoo source used for the academic pilot."""

    def __init__(self, base_url: str):
        self.base_url = base_url

    def fetch(self, symbol: str, start: date, end: date) -> PriceResult:
        yahoo_symbol = ticker_for_yahoo(symbol)
        period1 = int(datetime.combine(start, dt_time.min, timezone.utc).timestamp())
        period2 = int(datetime.combine(end, dt_time.max, timezone.utc).timestamp())
        url = self.base_url.format(symbol=urllib.parse.quote(yahoo_symbol))
        params = urllib.parse.urlencode(
            {
                "period1": period1,
                "period2": period2,
                "interval": "1d",
                "events": "div,splits",
                "includeAdjustedClose": "true",
            }
        )
        raw = _request_json(f"{url}?{params}")
        chart = raw.get("chart", {})
        if chart.get("error"):
            raise RuntimeError(f"Yahoo chart error for {symbol}: {chart['error']}")
        result = (chart.get("result") or [None])[0]
        if not result:
            raise RuntimeError(f"No price result for {symbol}")
        timestamps = result.get("timestamp") or []
        quote = ((result.get("indicators") or {}).get("quote") or [{}])[0]
        adjclose = ((result.get("indicators") or {}).get("adjclose") or [{}])[0].get(
            "adjclose", []
        )
        records = []
        for i, timestamp in enumerate(timestamps):
            def value(key: str) -> Any:
                values = quote.get(key) or []
                return values[i] if i < len(values) else None

            records.append(
                {
                    "date": datetime.fromtimestamp(timestamp, timezone.utc).date().isoformat(),
                    "open": value("open"),
                    "high": value("high"),
                    "low": value("low"),
                    "close": value("close"),
                    "adj_close": adjclose[i] if i < len(adjclose) else value("close"),
                    "volume": value("volume"),
                }
            )
        frame = pd.DataFrame.from_records(records)
        if frame.empty:
            raise RuntimeError(f"Empty price history for {symbol}")
        frame["date"] = pd.to_datetime(frame["date"])
        for column in ["open", "high", "low", "close", "adj_close", "volume"]:
            frame[column] = pd.to_numeric(frame[column], errors="coerce")
        frame = frame.dropna(subset=["adj_close"]).sort_values("date").reset_index(drop=True)
        metadata = dict(result.get("meta") or {})
        metadata.update(
            {
                "requested_symbol": symbol,
                "provider_symbol": yahoo_symbol,
                "source_url": url,
                "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
                "research_status": "PROTOTYPE_UNOFFICIAL_NO_KEY_SOURCE",
            }
        )
        return PriceResult(frame=frame, metadata=metadata, raw=raw, provider="Yahoo chart")


YAHOO_FUNDAMENTAL_METRICS = [
    "trailingMarketCap",
    "trailingPeRatio",
    "trailingTotalRevenue",
    "trailingNetIncomeCommonStockholders",
    "trailingDilutedEPS",
    "trailingFreeCashFlow",
    "trailingOperatingIncome",
    "trailingOperatingCashFlow",
    "quarterlyTotalDebt",
    "quarterlyStockholdersEquity",
    "quarterlyCashCashEquivalentsAndShortTermInvestments",
    "annualTotalRevenue",
    "annualNetIncomeCommonStockholders",
    "annualDilutedEPS",
    "annualFreeCashFlow",
]
YAHOO_METRIC_SET_SHA256 = hashlib.sha256(
    ",".join(YAHOO_FUNDAMENTAL_METRICS).encode("utf-8")
).hexdigest()

YAHOO_FUNDAMENTAL_OUTPUT_NAMES = {
    "trailingMarketCap": "market_cap",
    "trailingPeRatio": "trailing_pe",
    "trailingTotalRevenue": "revenue_ttm",
    "trailingNetIncomeCommonStockholders": "net_income_ttm",
    "trailingDilutedEPS": "diluted_eps_ttm",
    "trailingFreeCashFlow": "free_cash_flow_ttm",
    "trailingOperatingIncome": "operating_income_ttm",
    "trailingOperatingCashFlow": "operating_cash_flow_ttm",
    "quarterlyTotalDebt": "total_debt_latest_quarter",
    "quarterlyStockholdersEquity": "stockholders_equity_latest_quarter",
    "quarterlyCashCashEquivalentsAndShortTermInvestments": "cash_and_short_term_investments_latest_quarter",
}


def _json_safe(value: Any) -> Any:
    """Convert pandas/numpy values to stable JSON-native values."""
    if value is None:
        return None
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if hasattr(value, "item"):
        try:
            return _json_safe(value.item())
        except (TypeError, ValueError):
            pass
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def _frame_records(value: Any) -> list[dict[str, Any]]:
    if value is None:
        return []
    if isinstance(value, pd.DataFrame):
        frame = value.reset_index()
        return [_json_safe(record) for record in frame.to_dict(orient="records")]
    if isinstance(value, dict):
        return [_json_safe(value)]
    return []


def _current_recommendation(records: list[dict[str, Any]]) -> dict[str, Any]:
    if not records:
        return {}
    current = next(
        (
            record
            for record in records
            if str(record.get("period", "")).strip().lower() == "0m"
        ),
        None,
    )
    if current is None:
        return {}
    return {
        "period": str(current.get("period", "")),
        "strongBuy": int(current.get("strongBuy", 0) or 0),
        "buy": int(current.get("buy", 0) or 0),
        "hold": int(current.get("hold", 0) or 0),
        "sell": int(current.get("sell", 0) or 0),
        "strongSell": int(current.get("strongSell", 0) or 0),
    }


def _parse_fundamental_timeseries(
    document: dict[str, Any], cutoff: date
) -> dict[str, Any]:
    root = document.get("timeseries") or {}
    if root.get("error"):
        raise RuntimeError(f"Yahoo fundamentals error: {root['error']}")
    parsed: dict[str, Any] = {}
    for item in root.get("result") or []:
        type_value = (item.get("meta") or {}).get("type")
        metric = type_value[0] if isinstance(type_value, list) else type_value
        if not metric:
            continue
        observations = []
        for observation in item.get(metric) or []:
            raw_date = observation.get("asOfDate")
            try:
                observation_date = date.fromisoformat(str(raw_date))
            except (TypeError, ValueError):
                continue
            if (
                observation_date <= cutoff
                and (observation.get("reportedValue") or {}).get("raw") is not None
            ):
                observations.append(observation)
        observations.sort(key=lambda observation: str(observation.get("asOfDate", "")))
        parsed[str(metric)] = {
            "latest": _json_safe(observations[-1]) if observations else None,
            "history": _json_safe(observations),
        }
    available_metrics = {
        metric
        for metric in YAHOO_FUNDAMENTAL_METRICS
        if (parsed.get(metric) or {}).get("latest") is not None
    }
    parsed["missing_metrics"] = sorted(set(YAHOO_FUNDAMENTAL_METRICS) - available_metrics)
    if len(parsed["missing_metrics"]) == len(YAHOO_FUNDAMENTAL_METRICS):
        raise RuntimeError("Yahoo fundamentals response omitted every requested metric")
    return parsed


def _metric_packet(parsed: dict[str, Any], metric: str) -> dict[str, Any] | None:
    latest = (parsed.get(metric) or {}).get("latest")
    if not latest:
        return None
    return {
        "value": _json_safe((latest.get("reportedValue") or {}).get("raw")),
        "as_of_date": latest.get("asOfDate"),
        "period_type": latest.get("periodType"),
        "currency": latest.get("currencyCode"),
        "data_id": latest.get("dataId"),
    }


def _packet_value(packet: dict[str, Any], name: str) -> float | None:
    item = packet.get(name)
    if not isinstance(item, dict) or item.get("value") is None:
        return None
    try:
        return float(item["value"])
    except (TypeError, ValueError):
        return None


def _ratio(numerator: float | None, denominator: float | None) -> float | None:
    if numerator is None or denominator in (None, 0):
        return None
    return numerator / denominator


class YahooResearchProvider:
    """Fetch a compact rating/forecast-free fundamental packet and separate consensus.

    This adapter uses the third-party ``yfinance`` package over Yahoo's public
    interfaces. It is suitable for a small academic feasibility pilot, not a
    licensed production feed.
    """

    def __init__(self) -> None:
        try:
            import yfinance as yf
        except ImportError as exc:
            raise RuntimeError("Install requirements.txt before using Yahoo research data") from exc
        self.yf = yf

    def fetch(self, symbol: str, cutoff: date) -> ResearchResult:
        retrieval_date = datetime.now(timezone.utc).date()
        if cutoff != retrieval_date:
            raise ValueError(
                "Yahoo analyst data are current-only; fetch requires cutoff to equal "
                f"the UTC retrieval date ({retrieval_date.isoformat()})"
            )
        provider_symbol = ticker_for_yahoo(symbol)
        ticker = self.yf.Ticker(provider_symbol)
        retrieved_at = datetime.now(timezone.utc).isoformat()
        errors: dict[str, str] = {}

        start = datetime(cutoff.year - 6, 1, 1, tzinfo=timezone.utc)
        end = datetime.combine(cutoff + timedelta(days=1), dt_time.min, timezone.utc)
        params = urllib.parse.urlencode(
            {
                "symbol": provider_symbol,
                "type": ",".join(YAHOO_FUNDAMENTAL_METRICS),
                "merge": "false",
                "period1": str(int(start.timestamp())),
                "period2": str(int(end.timestamp())),
                "lang": "en-US",
                "region": "US",
            }
        )
        fundamentals_url = (
            "https://query1.finance.yahoo.com/ws/fundamentals-timeseries/v1/"
            f"finance/timeseries/{urllib.parse.quote(provider_symbol)}?{params}"
        )
        try:
            fundamentals_raw = _request_json(fundamentals_url)
            parsed_fundamentals = _parse_fundamental_timeseries(fundamentals_raw, cutoff)
        except Exception as exc:  # Yahoo can omit modules or rate-limit individual symbols.
            fundamentals_raw = {}
            parsed_fundamentals = {"missing_metrics": list(YAHOO_FUNDAMENTAL_METRICS)}
            errors["fundamentals_timeseries"] = f"{type(exc).__name__}: {exc}"

        try:
            info = ticker.get_info() or {}
        except Exception as exc:
            info = {}
            errors["portal_summary"] = f"{type(exc).__name__}: {exc}"

        try:
            recommendations = ticker.get_recommendations()
        except Exception as exc:
            recommendations = None
            errors["recommendations"] = f"{type(exc).__name__}: {exc}"

        try:
            price_targets = ticker.get_analyst_price_targets() or {}
        except Exception as exc:
            price_targets = {}
            errors["analyst_price_targets"] = f"{type(exc).__name__}: {exc}"

        recommendation_records = _frame_records(recommendations)
        consensus = _current_recommendation(recommendation_records)
        yahoo_sector = _json_safe(info.get("sector"))
        gics_sector = normalize_yahoo_sector(yahoo_sector)
        consensus.update(
            {
                "portalRecommendationKey": info.get("recommendationKey"),
                "portalRecommendationMean": _json_safe(info.get("recommendationMean")),
                "portalNumberOfAnalystOpinions": _json_safe(
                    info.get("numberOfAnalystOpinions")
                ),
                "priceTargetCurrent": _json_safe(price_targets.get("current")),
                "priceTargetLow": _json_safe(price_targets.get("low")),
                "priceTargetMean": _json_safe(price_targets.get("mean")),
                "priceTargetMedian": _json_safe(price_targets.get("median")),
                "priceTargetHigh": _json_safe(price_targets.get("high")),
                "retrievedAtUtc": retrieved_at,
            }
        )

        fundamentals = {}
        for metric, output_name in YAHOO_FUNDAMENTAL_OUTPUT_NAMES.items():
            packet = _metric_packet(parsed_fundamentals, metric)
            if packet is not None:
                fundamentals[output_name] = packet

        revenue = _packet_value(fundamentals, "revenue_ttm")
        net_income = _packet_value(fundamentals, "net_income_ttm")
        operating_income = _packet_value(fundamentals, "operating_income_ttm")
        free_cash_flow = _packet_value(fundamentals, "free_cash_flow_ttm")
        total_debt = _packet_value(fundamentals, "total_debt_latest_quarter")
        equity = _packet_value(fundamentals, "stockholders_equity_latest_quarter")
        cash = _packet_value(
            fundamentals, "cash_and_short_term_investments_latest_quarter"
        )
        fundamentals["locally_derived"] = {
            "operating_margin_ttm": _ratio(operating_income, revenue),
            "net_margin_ttm": _ratio(net_income, revenue),
            "free_cash_flow_margin_ttm": _ratio(free_cash_flow, revenue),
            "debt_to_equity_latest_quarter": _ratio(total_debt, equity),
            "net_cash_latest_quarter": (
                cash - total_debt if cash is not None and total_debt is not None else None
            ),
        }
        fundamentals.update(
            {
                "missing_metrics": parsed_fundamentals.get("missing_metrics", []),
                "data_retrieved_at_utc": retrieved_at,
                "data_cutoff": cutoff.isoformat(),
                "source": "Yahoo Finance fundamentals-timeseries (unofficial)",
            }
        )
        metadata = {
            "provider": "Yahoo Finance via yfinance",
            "schema_version": YAHOO_RESEARCH_SCHEMA_VERSION,
            "metric_set_sha256": YAHOO_METRIC_SET_SHA256,
            "provider_symbol": provider_symbol,
            "sector_yahoo": yahoo_sector,
            "sector_gics": gics_sector,
            "retrieved_at_utc": retrieved_at,
            "data_cutoff": cutoff.isoformat(),
            "quote_url": f"https://finance.yahoo.com/quote/{urllib.parse.quote(provider_symbol)}/",
            "analyst_url": (
                f"https://finance.yahoo.com/quote/{urllib.parse.quote(provider_symbol)}/"
                "analyst-insights/"
            ),
            "fundamentals_url": fundamentals_url,
            "yfinance_version": getattr(self.yf, "__version__", None),
            "research_status": "UNOFFICIAL_ACADEMIC_PILOT_SOURCE",
            "errors": errors,
        }
        raw = {
            "portal_summary": {
                "recommendationKey": _json_safe(info.get("recommendationKey")),
                "recommendationMean": _json_safe(info.get("recommendationMean")),
                "numberOfAnalystOpinions": _json_safe(info.get("numberOfAnalystOpinions")),
                "sector": yahoo_sector,
                "sector_gics": gics_sector,
            },
            "fundamentals_timeseries": _json_safe(fundamentals_raw),
            "recommendations": recommendation_records,
            "analyst_price_targets": _json_safe(price_targets),
            "errors": errors,
        }
        return ResearchResult(
            fundamentals=fundamentals,
            analyst_consensus=consensus,
            metadata=metadata,
            raw=raw,
            provider="Yahoo Finance via yfinance",
        )


class FMPProvider:
    """Production adapter for Financial Modeling Prep stable endpoints."""

    def __init__(self, api_key: str, base_url: str):
        if not api_key:
            raise ValueError("FMP_API_KEY is required")
        self.api_key = api_key
        self.base_url = base_url.rstrip("/")

    def _get(self, endpoint: str, **params: Any) -> Any:
        params["apikey"] = self.api_key
        query = urllib.parse.urlencode(params)
        return _request_json(f"{self.base_url}/{endpoint}?{query}")

    def fetch_universe(self) -> pd.DataFrame:
        data = self._get("sp500-constituent")
        frame = pd.DataFrame(data)
        frame = frame.rename(columns={"name": "company", "companyName": "company", "subSector": "sub_sector"})
        required = {"symbol", "company", "sector"}
        missing = required - set(frame.columns)
        if missing:
            raise RuntimeError(f"FMP constituent response is missing columns: {sorted(missing)}")
        return frame

    def fetch_prices(self, symbol: str, start: date, end: date) -> PriceResult:
        data = self._get(
            "historical-price-eod/dividend-adjusted",
            symbol=symbol,
            **{"from": start.isoformat(), "to": end.isoformat()},
        )
        records = data.get("historical", data) if isinstance(data, dict) else data
        frame = pd.DataFrame(records)
        if frame.empty:
            raise RuntimeError(f"FMP returned no price history for {symbol}")
        frame = frame.rename(columns={"adjClose": "adj_close"})
        if "adj_close" not in frame:
            frame["adj_close"] = frame["close"]
        frame["date"] = pd.to_datetime(frame["date"])
        frame = frame.sort_values("date").reset_index(drop=True)
        metadata = {
            "requested_symbol": symbol,
            "source_endpoint": "historical-price-eod/dividend-adjusted",
            "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
        }
        return PriceResult(frame=frame, metadata=metadata, raw={"data": data}, provider="FMP")

    def fetch_fundamentals(self, symbol: str) -> dict[str, Any]:
        metrics = self._get("key-metrics-ttm", symbol=symbol)
        ratios = self._get("ratios-ttm", symbol=symbol)
        first_metrics = metrics[0] if isinstance(metrics, list) and metrics else {}
        first_ratios = ratios[0] if isinstance(ratios, list) and ratios else {}
        return {**first_metrics, **first_ratios}

    def fetch_analyst_consensus(self, symbol: str) -> dict[str, Any]:
        data = self._get("grades-consensus", symbol=symbol)
        return data[0] if isinstance(data, list) and data else (data or {})
